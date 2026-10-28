# -*- coding: utf-8 -*-
"""Genera, para la seccion Reporteria de la web:
  - <ID>.json por cada Excel de artefacto (columnas + filas de la tabla real)
    para mostrarlo como listado/grid en el cliente.
  - artefactos.zip con todos los Excel + el diagrama (descarga "todos de frente").
Lee los .xlsx ya copiados en client/public/artefactos/.
Uso: python scripts/gen_artefactos.py
"""
import json, zipfile
from pathlib import Path
from openpyxl import load_workbook

PUB = Path(__file__).resolve().parent.parent / "public"
ART = PUB / "artefactos"

def extract(xlsx):
    wb = load_workbook(xlsx, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    # titulo: primera celda que contenga 'ADI-'
    titulo = xlsx.stem
    for r in rows[:6]:
        for v in r:
            if v and "ADI-" in str(v):
                titulo = str(v).strip(); break
        if titulo != xlsx.stem:
            break
    # fila de encabezado = la de mas celdas no vacias en las primeras 25
    best_i, best_n = 0, 0
    for i, r in enumerate(rows[:25]):
        n = sum(1 for v in r if v is not None and str(v).strip())
        if n > best_n:
            best_n, best_i = n, i
    hdr = rows[best_i]
    # indices de columnas no vacias del encabezado
    idx = [j for j, v in enumerate(hdr) if v is not None and str(v).strip()]
    columns = [str(hdr[j]).strip() for j in idx]
    data = []
    for r in rows[best_i + 1:]:
        cells = [r[j] if j < len(r) else None for j in idx]
        if all(c is None or str(c).strip() == "" for c in cells):
            if data:
                break
            continue
        data.append(["" if c is None else str(c) for c in cells])
    return {"id": xlsx.stem, "titulo": titulo, "columns": columns, "rows": data}

def main():
    xlsxs = sorted(ART.glob("ADI-*.xlsx"))
    for x in xlsxs:
        d = extract(x)
        (ART / f"{x.stem}.json").write_text(json.dumps(d, ensure_ascii=False), encoding="utf-8")
        print(f"  {x.stem}.json  -> {len(d['columns'])} cols x {len(d['rows'])} filas  ({d['titulo'][:50]})")
    # zip con todos los Excel + el diagrama
    zpath = ART / "artefactos.zip"
    with zipfile.ZipFile(zpath, "w", zipfile.ZIP_DEFLATED) as z:
        for x in xlsxs:
            z.write(x, x.name)
        flujo = PUB / "flujo_datos_onp.html"
        if flujo.exists():
            z.write(flujo, flujo.name)
    print(f"artefactos.zip -> {zpath.stat().st_size} bytes ({len(xlsxs)} xlsx + diagrama)")

if __name__ == "__main__":
    main()
